"""Write-back — two outputs from one pricing run (CLAUDE.md §7).

a. CLIENT copy: the client's original BOQ file with Rate/Amount landed as
   VALUES ONLY into the detected cells. Structure/format untouched: we edit
   the original workbook object in place and never touch any other cell.
   Vendor base rates, markup, and formulas are never exposed.
b. INTERNAL workbook: live formulas — every rate references the markup cell
   (=ROUND(base*(1+$Markup$),r)), plus a vendor comparison sheet and a
   margin table per item/package.

Metadata hygiene (INVENTORY hazard #16): both outputs get their docProps
creator/lastModifiedBy scrubbed so internal names never ship.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import openpyxl

from .price import PricedItem, PricingRun


def scrub_metadata(path: str | Path, author: str = "Estimation") -> None:
    """Rewrite docProps in place (zip-level, nothing else changes)."""
    import os
    import re

    path = str(path)
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "docProps/core.xml":
                s = data.decode("utf-8")
                s = re.sub(r"<dc:creator>[^<]*</dc:creator>", f"<dc:creator>{author}</dc:creator>", s)
                s = re.sub(r"<cp:lastModifiedBy>[^<]*</cp:lastModifiedBy>",
                           f"<cp:lastModifiedBy>{author}</cp:lastModifiedBy>", s)
                data = s.encode("utf-8")
            zout.writestr(info, data)
    os.replace(tmp, path)


def write_client_copy(original_boq: str | Path, priced: list[PricedItem],
                      out_path: str | Path, rate_col_by_sheet: dict[str, dict]) -> dict:
    """Land final rate/amount as values into the original workbook.

    rate_col_by_sheet: sheet -> {"rate": col_idx, "amount": col_idx|None,
    "supply_rate": ..., "install_rate": ...} from the extraction HeaderMap —
    values go ONLY into those detected cells on the priced rows.
    """
    wb = openpyxl.load_workbook(original_boq)  # keep formulas/format of everything else
    touched = 0
    skipped = []
    for it in priced:
        if it.status != "priced" or it.final_rate is None:
            continue
        cols = rate_col_by_sheet.get(it.sheet)
        if not cols or it.sheet not in wb.sheetnames:
            skipped.append({"sheet": it.sheet, "row": it.row, "reason": "no rate column mapping"})
            continue
        ws = wb[it.sheet]
        rate_c = cols.get("rate") or cols.get("supply_rate")
        if rate_c is None:
            skipped.append({"sheet": it.sheet, "row": it.row, "reason": "no rate column"})
            continue
        ws.cell(row=it.row, column=rate_c).value = it.final_rate
        # amount: value only — never a formula in the client copy
        amt_c = cols.get("amount") or cols.get("supply_amount")
        if amt_c is not None and it.amount is not None:
            ws.cell(row=it.row, column=amt_c).value = it.amount
        touched += 1
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    scrub_metadata(out_path)
    return {"file": str(out_path), "cells_filled": touched, "skipped": skipped}


def write_internal_workbook(run: PricingRun, out_path: str | Path,
                            project: str = "project") -> dict:
    """Internal working workbook with live formulas.

    Sheet 'Control': markup cell (B2) referenced by every rate formula, GST
    mode/pct, area. Sheet 'Pricing': one row per item with
    =ROUND(base*(1+Markup),r) rates. Sheet 'Comparison': all vendor
    candidates per item, L1 highlighted. Sheet 'Margin': per-package rollup
    with live SUMIF formulas.
    """
    wb = openpyxl.Workbook()
    ctl = wb.active
    ctl.title = "Control"
    ctl["A1"] = "Setting"
    ctl["B1"] = "Value"
    ctl["A2"] = "Markup %"
    ctl["B2"] = run.markup_pct
    ctl["A3"] = "GST mode"
    ctl["B3"] = run.gst_mode
    ctl["A4"] = "GST %"
    ctl["B4"] = run.gst_pct
    ctl["A5"] = "Area (sqft)"
    ctl["B5"] = run.area_sqft
    ctl["A6"] = "Round digits"
    ctl["B6"] = run.round_digits

    pr = wb.create_sheet("Pricing")
    headers = ["Package", "Sheet", "Row", "Description", "Unit", "Qty",
               "L1 Vendor", "L1 Rate", "Final Rate", "Amount", "Cost", "Margin", "Status"]
    pr.append(headers)
    r = 2
    for it in run.items:
        base = it.l1.total_rate if it.l1 else None
        pr.append([
            it.package, it.sheet, it.row, it.description[:200], it.unit, it.qty,
            it.l1.vendor if it.l1 else None, base,
            (f"=ROUND(H{r}*(1+Control!$B$2/100),{run.round_digits})" if base is not None else None),
            (f"=ROUND(F{r}*I{r},2)" if base is not None and it.qty else None),
            (f"=ROUND(F{r}*H{r},2)" if base is not None and it.qty else None),
            (f"=J{r}-K{r}" if base is not None and it.qty else None),
            it.status,
        ])
        r += 1
    last = r - 1

    cmp_ws = wb.create_sheet("Comparison")
    cmp_ws.append(["Sheet", "Row", "Description", "Vendor", "Total Rate",
                   "Supply", "Install", "Confidence", "Signal", "Is L1", "History flag"])
    for it in run.items:
        for c in sorted(it.candidates, key=lambda x: x.total_rate or 9e18):
            cmp_ws.append([
                it.sheet, it.row, it.description[:120], c.vendor, c.total_rate,
                c.supply_rate, c.install_rate, c.confidence, c.signal,
                "L1" if it.l1 is c else "",
                str(c.history_flag) if c.history_flag else "",
            ])

    mg = wb.create_sheet("Margin")
    mg.append(["Package", "Offer Amount", "Cost", "Margin", "Margin %", "₹/sqft"])
    packages = sorted({it.package for it in run.items})
    mr = 2
    for pkg in packages:
        mg.append([
            pkg,
            f'=SUMIF(Pricing!A2:A{last},A{mr},Pricing!J2:J{last})',
            f'=SUMIF(Pricing!A2:A{last},A{mr},Pricing!K2:K{last})',
            f"=B{mr}-C{mr}",
            f"=IF(B{mr}=0,\"\",D{mr}/B{mr}*100)",
            (f"=IF(Control!$B$5>0,B{mr}/Control!$B$5,\"\")" if run.area_sqft else ""),
        ])
        mr += 1
    mg.append([
        "TOTAL (ex GST)", f"=SUM(B2:B{mr - 1})", f"=SUM(C2:C{mr - 1})",
        f"=B{mr}-C{mr}", f"=IF(B{mr}=0,\"\",D{mr}/B{mr}*100)",
        (f"=IF(Control!$B$5>0,B{mr}/Control!$B$5,\"\")" if run.area_sqft else ""),
    ])
    mg.append(["GST", f"=IF(Control!$B$3=\"exclusive\",B{mr}*Control!$B$4/100,0)"])
    mg.append(["GRAND TOTAL", f"=B{mr}+B{mr + 1}"])

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    scrub_metadata(out_path)
    return {"file": str(out_path), "items": len(run.items), "packages": len(packages)}
