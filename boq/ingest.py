"""Load xlsx/csv workbooks into coordinate-addressed grids.

Design constraints from fixtures/INVENTORY.md:
- merged cells must be resolved (top-left value propagated) — hazard #3
- #REF!/external-link formulas must degrade gracefully — hazard #1
- ArrayFormula objects need their own handling — hazard #2
- lock files (~$...) and junk must be skipped — hazard #15
- cached values are lost after any openpyxl re-save — hazard #17, so we read
  BOTH the formula view and the cached-value view in one pass and keep them.
"""

from __future__ import annotations

import csv
import hashlib
import json
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

JUNK_PREFIXES = ("~$",)
JUNK_NAMES = {".DS_Store", "Thumbs.db"}
EXCEL_EXTS = {".xlsx", ".xlsm"}
CSV_EXTS = {".csv"}


@dataclass
class Cell:
    """One cell: raw formula/value + cached value + coordinates."""

    row: int
    col: int
    value: object  # literal value, or cached value when the cell is a formula
    formula: str | None = None  # "=..." when the cell holds a formula
    is_array_formula: bool = False

    @property
    def coord(self) -> str:
        return f"{get_column_letter(self.col)}{self.row}"

    @property
    def text(self) -> str:
        """Best-effort human text for matching/detection."""
        if self.value is None:
            return ""
        return str(self.value).strip()


@dataclass
class SheetGrid:
    name: str
    cells: dict[tuple[int, int], Cell] = field(default_factory=dict)
    merged_ranges: list[str] = field(default_factory=list)
    max_row: int = 0
    max_col: int = 0

    def cell(self, row: int, col: int) -> Cell | None:
        return self.cells.get((row, col))

    def value(self, row: int, col: int) -> object:
        c = self.cells.get((row, col))
        return c.value if c else None

    def text(self, row: int, col: int) -> str:
        c = self.cells.get((row, col))
        return c.text if c else ""

    def row_texts(self, row: int, max_col: int | None = None) -> list[str]:
        n = max_col or self.max_col
        return [self.text(row, c) for c in range(1, n + 1)]

    def iter_rows(self):
        for r in range(1, self.max_row + 1):
            yield r, [self.cells.get((r, c)) for c in range(1, self.max_col + 1)]


@dataclass
class Workbook:
    path: Path
    sheets: dict[str, SheetGrid] = field(default_factory=dict)
    file_hash: str = ""

    @property
    def sheet_names(self) -> list[str]:
        return list(self.sheets)


def is_junk(path: Path) -> bool:
    return path.name.startswith(JUNK_PREFIXES) or path.name in JUNK_NAMES


def file_hash(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _formula_str(v: object) -> str | None:
    if isinstance(v, str) and v.startswith("="):
        return v
    if isinstance(v, ArrayFormula):
        return v.text or "="
    return None


def load_workbook(path: str | Path, sheet_filter=None) -> Workbook:
    """Load an Excel workbook into grids, resolving merges.

    Two openpyxl passes: data_only=False for formulas, data_only=True for the
    cached values Excel last computed. On files that were re-saved by openpyxl
    the cache is empty — cells then carry value=None for formula cells, which
    downstream code must treat as "value unknown", not zero (hazard #17).
    """
    path = Path(path)
    if is_junk(path):
        raise ValueError(f"junk/lock file, refusing to parse: {path.name}")
    if path.suffix.lower() in CSV_EXTS:
        return _load_csv(path)
    if path.suffix.lower() not in EXCEL_EXTS:
        raise ValueError(f"unsupported file type: {path.suffix} ({path.name})")

    wb_f = openpyxl.load_workbook(path, data_only=False, read_only=False)
    wb_v = openpyxl.load_workbook(path, data_only=True, read_only=False)
    out = Workbook(path=path, file_hash=file_hash(path))

    for sn in wb_f.sheetnames:
        if sheet_filter and not sheet_filter(sn):
            continue
        ws_f, ws_v = wb_f[sn], wb_v[sn]
        grid = SheetGrid(name=sn, max_row=ws_f.max_row or 0, max_col=ws_f.max_column or 0)
        grid.merged_ranges = [str(r) for r in ws_f.merged_cells.ranges]

        for row in ws_f.iter_rows():
            for cf in row:
                if cf.value is None:
                    continue
                formula = _formula_str(cf.value)
                cached = ws_v.cell(row=cf.row, column=cf.column).value if formula else cf.value
                # "#REF!" and friends come back as literal strings in the value
                # view; keep them (extract.py treats them as missing metadata).
                grid.cells[(cf.row, cf.column)] = Cell(
                    row=cf.row,
                    col=cf.column,
                    value=cached,
                    formula=formula,
                    is_array_formula=isinstance(cf.value, ArrayFormula),
                )

        # resolve merges: propagate top-left cell into every covered position
        for rng in ws_f.merged_cells.ranges:
            tl = grid.cells.get((rng.min_row, rng.min_col))
            if tl is None:
                continue
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    if (r, c) not in grid.cells:
                        grid.cells[(r, c)] = Cell(row=r, col=c, value=tl.value, formula=tl.formula)
        out.sheets[sn] = grid

    wb_f.close()
    wb_v.close()
    return out


def _load_csv(path: Path) -> Workbook:
    out = Workbook(path=path, file_hash=file_hash(path))
    grid = SheetGrid(name=path.stem)
    with open(path, newline="", encoding="utf-8-sig") as f:
        for r, row in enumerate(csv.reader(f), start=1):
            for c, val in enumerate(row, start=1):
                if val != "":
                    grid.cells[(r, c)] = Cell(row=r, col=c, value=val)
                    grid.max_row = max(grid.max_row, r)
                    grid.max_col = max(grid.max_col, c)
    out.sheets[grid.name] = grid
    return out


def grid_to_text(grid: SheetGrid, max_rows: int = 400, max_chars_per_cell: int = 300) -> str:
    """Coordinate-addressed text dump of a sheet (for LLM prompts and debugging)."""
    lines = [f"# sheet: {grid.name} ({grid.max_row}x{grid.max_col}, merged: {len(grid.merged_ranges)})"]
    for r in range(1, min(grid.max_row, max_rows) + 1):
        parts = []
        for c in range(1, grid.max_col + 1):
            cell = grid.cell(r, c)
            if cell is None or cell.text == "":
                continue
            t = cell.text[:max_chars_per_cell].replace("\n", " ")
            parts.append(f"{cell.coord}={t}")
        if parts:
            lines.append(" | ".join(parts))
    return "\n".join(lines)


def cache_extraction(cache_dir: str | Path, wb: Workbook, payload: dict) -> Path:
    """File-hash-keyed cache of extraction results."""
    cache_dir = Path(cache_dir)
    cache_dir.mkdir(parents=True, exist_ok=True)
    p = cache_dir / f"{wb.file_hash}.json"
    p.write_text(json.dumps(payload, indent=1, default=str), encoding="utf-8")
    return p


def load_cached(cache_dir: str | Path, path: str | Path) -> dict | None:
    p = Path(cache_dir) / f"{file_hash(Path(path))}.json"
    if p.exists():
        return json.loads(p.read_text(encoding="utf-8"))
    return None
