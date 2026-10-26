"""Float packages (CLAUDE.md §3, M7).

Group BOQ items by trade and generate per-vendor float packages: an Excel
extract of that vendor's items with BLANK rate columns and make constraints
noted from the make list. Mirrors the team's real process: category ->
vendors, tracked with SENT/RECEIVED status (float-tracker fixture schema).
"""

from __future__ import annotations

import json
import re
from datetime import date
from pathlib import Path

import openpyxl

from .extract import extract_workbook
from .ingest import load_workbook
from .knowledge import fingerprint_tokens
from .writeback import scrub_metadata

# trade grouping: sheet-name driven first (the tender's own sheet split is the
# house trade taxonomy), keyword fallback for single-sheet BOQs
TRADE_KEYWORDS = {
    "flooring": re.compile(r"floor|tile|carpet|lvt|vinyl", re.I),
    "ceiling": re.compile(r"ceiling|gypsum|baffle", re.I),
    "partitions & doors": re.compile(r"partition|door|glazing", re.I),
    "electrical": re.compile(r"electrical|cable|db|panel|wiring|light", re.I),
    "hvac": re.compile(r"hvac|ac\b|vrf|duct|ahu|odu", re.I),
    "plumbing": re.compile(r"plumb|sanitary|cp fitting|wc|urinal", re.I),
    "furniture": re.compile(r"furniture|chair|table|sofa|locker", re.I),
    "painting & finishes": re.compile(r"paint|polish|wallpaper|finish", re.I),
}


def infer_trade(sheet_name: str, description: str = "") -> str:
    for trade, pat in TRADE_KEYWORDS.items():
        if pat.search(sheet_name):
            return trade
    for trade, pat in TRADE_KEYWORDS.items():
        if description and pat.search(description):
            return trade
    return sheet_name.strip() or "general"


def make_constraints(make_list_file: str | Path) -> dict[str, list[str]]:
    """sheet/trade -> approved makes, from the make-list workbook."""
    payload = extract_workbook(load_workbook(make_list_file))
    out: dict[str, list[str]] = {}
    for s in payload["sheets"]:
        makes = sorted({
            (it.get("make") or "").strip()
            for it in s["items"] if (it.get("make") or "").strip()
        })
        if makes:
            out[s["sheet"].strip()] = makes
    return out


def match_constraint_sheet(boq_sheet: str, constraints: dict[str, list[str]]) -> list[str]:
    """Make-list tabs rarely equal BOQ tabs exactly — token match them."""
    b = fingerprint_tokens(boq_sheet)
    best, score = [], 0.0
    for name, makes in constraints.items():
        n = fingerprint_tokens(name)
        if not n or not b:
            continue
        sc = len(b & n) / max(len(b | n), 1)
        if sc > score:
            best, score = makes, sc
    return best if score >= 0.3 else []


def build_float_packages(boq_file: str | Path, make_list: str | Path | None = None,
                         vendor_map: dict[str, list[str]] | None = None,
                         out_dir: str | Path = "floats") -> dict:
    """One xlsx per (trade, vendor) — or per trade when no vendor map given.

    vendor_map: trade -> [vendor names] (the tracker's category->vendor rows).
    """
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    payload = extract_workbook(load_workbook(boq_file))
    constraints = make_constraints(make_list) if make_list else {}

    trades: dict[str, list[dict]] = {}
    for s in payload["sheets"]:
        if s["kind"] != "boq":
            continue
        for it in s["items"]:
            if not (it.get("description") or "").strip():
                continue
            trade = infer_trade(s["sheet"], it.get("description", ""))
            trades.setdefault(trade, []).append(it)

    tracker_rows = []
    files = []
    for trade, items in sorted(trades.items()):
        vendors = (vendor_map or {}).get(trade, [None])
        for vendor in vendors:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Float"
            ws["A1"] = "FLOAT PACKAGE"
            ws["B1"] = trade.upper()
            ws["A2"] = "Vendor"
            ws["B2"] = vendor or "<vendor>"
            ws["A3"] = "Date"
            ws["B3"] = date.today().isoformat()
            makes = match_constraint_sheet(items[0]["sheet"], constraints) if items else []
            ws["A4"] = "Approved makes"
            ws["B4"] = ", ".join(makes) if makes else "as per make list"
            ws.append([])
            ws.append(["Sl. No", "Source Sheet", "Description", "Make constraint",
                       "Unit", "Qty", "Supply Rate", "Installation Rate",
                       "Total Rate", "Amount", "Remarks"])
            hdr_row = ws.max_row
            for i, it in enumerate(items, start=1):
                desc = (it.get("description") or "")
                if it.get("spec"):
                    desc = desc + "\n" + it["spec"]
                r = ws.max_row + 1
                ws.append([
                    i, it["sheet"], desc[:2000],
                    it.get("make") or "", it.get("unit") or "", it.get("qty"),
                    None, None,  # rate columns intentionally blank
                    f"=IF(AND(G{r}=\"\",H{r}=\"\"),\"\",N(G{r})+N(H{r}))",
                    f"=IF(I{r}=\"\",\"\",ROUND(F{r}*I{r},2))",
                    "",
                ])
            # basic usability formatting
            ws.column_dimensions["C"].width = 80
            for col in "ABDEFGHIJK":
                ws.column_dimensions[col].width = 14
            fname = re.sub(r"[^\w\-]+", "-", trade.lower()).strip("-")
            if vendor:
                fname += "--" + re.sub(r"[^\w\-]+", "-", vendor.lower()).strip("-")
            path = out_dir / f"float-{fname}.xlsx"
            wb.save(path)
            scrub_metadata(path)
            files.append(str(path))
            tracker_rows.append({"trade": trade, "vendor": vendor,
                                 "items": len(items), "file": str(path),
                                 "status": "DRAFT"})

    tracker = {"generated": date.today().isoformat(),
               "boq_file": str(boq_file), "rows": tracker_rows}
    (out_dir / "float-tracker.json").write_text(json.dumps(tracker, indent=1), encoding="utf-8")
    return {"packages": len(files), "trades": len(trades),
            "files": files, "tracker": str(out_dir / "float-tracker.json")}
