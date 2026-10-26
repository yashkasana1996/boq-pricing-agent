"""Generate a fully synthetic demo tender pack.

Every company, vendor, item and rate here is invented. The files mimic the
shapes real fit-out tenders arrive in (multi-section BOQ, a quote that splits
supply and installation, a quote with amounts only, an in-house rate list)
so the pipeline can be tried end to end without any confidential data.

    python examples/make_demo_tender.py      # writes examples/demo-tender/
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

OUT = Path(__file__).resolve().parent / "demo-tender"

THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
SEC_FILL = PatternFill("solid", fgColor="D9E1F2")

# (sr, description, unit, qty) grouped by section
BOQ = {
    "A. FLOORING": [
        ("1", "Supply and laying of 600x600 mm vitrified tiles, 10 mm thick, in cement mortar", "Sqm", 820),
        ("2", "Supply and fixing of 100 mm high vitrified tile skirting matching floor tiles", "Rmt", 410),
        ("3", "Supply and laying of 6 mm thick carpet tiles 500x500 mm with underlay", "Sqm", 1150),
    ],
    "B. FALSE CEILING": [
        ("4", "Gypsum board false ceiling 12.5 mm thick on GI framework, painted", "Sqm", 1300),
        ("5", "Mineral fibre grid ceiling tiles 600x600 mm, 15 mm thick, on T-grid", "Sqm", 640),
    ],
    "C. PARTITIONS": [
        ("6", "Full height glass partition, 12 mm toughened glass with aluminium frame", "Sqm", 210),
        ("7", "Gypsum board partition 2x12.5 mm both sides with 50 mm rock wool infill", "Sqm", 380),
    ],
    "D. PAINTING": [
        ("8", "Two coats of acrylic emulsion paint over one coat primer on walls", "Sqm", 2600),
    ],
}

# vendor quotes: sr -> rates (invented). Vendor A splits supply/install.
VENDOR_A = {  # flooring + partitions specialist, supply + install
    "1": (1450.0, 380.0, "600x600 mm vitrified tiles 10 mm laid in cement mortar"),
    "2": (310.0, 95.0, "Vitrified tile skirting 100 mm high"),
    "3": (1680.0, 210.0, "Carpet tiles 500x500 mm 6 mm thick incl. underlay"),
    "6": (4900.0, 950.0, "12 mm toughened glass partition with aluminium frame"),
}
VENDOR_B = {  # ceilings + partitions, total rate only, amount column
    "4": (1125.0, "Gypsum board ceiling 12.5 mm on GI frame, painted"),
    "5": (980.0, "Mineral fibre tiles 600x600x15 mm on T-grid"),
    "6": (5620.0, "Toughened glass partition 12 mm, aluminium frame"),
    "7": (2350.0, "Gypsum partition 2x12.5 mm both sides, 50 mm rock wool"),
}
IN_HOUSE = {  # internal rate list used like one more vendor
    "1": (1890.0, "Vitrified tiles 600x600 10 mm"),
    "8": (165.0, "Acrylic emulsion 2 coats over primer"),
    "4": (1190.0, "Gypsum false ceiling 12.5 mm painted"),
}


def _header(ws, row, labels, widths=None):
    for c, label in enumerate(labels, start=1):
        cell = ws.cell(row=row, column=c, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.border = BOX
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for c, w in enumerate(widths or [], start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w


def make_boq(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOQ"
    ws["A1"] = "Northwind Corporate Office Fit-out (DEMO)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Bill of Quantities - Civil & Interior Works - Revision R0"
    _header(ws, 4, ["Sr. No.", "Description", "Unit", "Qty", "Rate (INR)", "Amount (INR)"],
            [8, 70, 8, 10, 14, 16])
    r = 5
    for section, items in BOQ.items():
        ws.cell(row=r, column=2, value=section).font = Font(bold=True)
        for c in range(1, 7):
            ws.cell(row=r, column=c).fill = SEC_FILL
        r += 1
        for sr, desc, unit, qty in items:
            ws.cell(row=r, column=1, value=sr)
            ws.cell(row=r, column=2, value=desc).alignment = Alignment(wrap_text=True)
            ws.cell(row=r, column=3, value=unit)
            ws.cell(row=r, column=4, value=qty)
            ws.cell(row=r, column=6, value=f"=D{r}*E{r}")
            for c in range(1, 7):
                ws.cell(row=r, column=c).border = BOX
            r += 1
    wb.save(path)


def make_vendor_a(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quotation"
    ws["A1"] = "Contoso Floors & Glass Pvt Ltd (DEMO) - Quotation"
    _header(ws, 3, ["BOQ Sr", "Description", "Unit", "Qty", "Supply Rate", "Installation Rate",
                    "Total Rate", "Amount"], [8, 60, 8, 10, 12, 14, 12, 14])
    qty = {sr: q for items in BOQ.values() for sr, _, _, q in items}
    unit = {sr: u for items in BOQ.values() for sr, _, u, _ in items}
    r = 4
    for sr, (s, i, desc) in VENDOR_A.items():
        ws.append([sr, desc, unit[sr], qty[sr], s, i, s + i, round((s + i) * qty[sr], 2)])
        r += 1
    ws.cell(row=r + 1, column=2, value="Terms: GST extra at actuals. Validity 30 days.")
    wb.save(path)


def make_vendor_b(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Offer"
    ws["A1"] = "Fabrikam Interiors LLP (DEMO) - Commercial Offer"
    _header(ws, 3, ["Sl No", "Item Description", "UOM", "Quantity", "Rate", "Amount"],
            [8, 60, 8, 10, 12, 14])
    qty = {sr: q for items in BOQ.values() for sr, _, _, q in items}
    unit = {sr: u for items in BOQ.values() for sr, _, u, _ in items}
    for sr, (rate, desc) in VENDOR_B.items():
        ws.append([sr, desc, unit[sr], qty[sr], rate, round(rate * qty[sr], 2)])
    wb.save(path)


def make_in_house(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rate List"
    _header(ws, 1, ["Ref", "Description", "Unit", "Rate"], [8, 50, 8, 12])
    unit = {sr: u for items in BOQ.values() for sr, _, u, _ in items}
    for sr, (rate, desc) in IN_HOUSE.items():
        ws.append([sr, desc, unit[sr], rate])
    wb.save(path)


def make_boq_r1(path: Path) -> None:
    """Client re-issue: carpet qty cut, mineral-fibre ceiling descoped,
    a new acoustic panel item added."""
    global BOQ
    original = {k: list(v) for k, v in BOQ.items()}
    BOQ["A. FLOORING"][2] = ("3", BOQ["A. FLOORING"][2][1], "Sqm", 900)
    BOQ["B. FALSE CEILING"] = [BOQ["B. FALSE CEILING"][0]]
    BOQ["D. PAINTING"] = BOQ["D. PAINTING"] + [
        ("9", "Fabric wrapped acoustic wall panels 25 mm thick with 50 mm air gap", "Sqm", 120)]
    try:
        make_boq(path)
    finally:
        BOQ = original


def main() -> None:
    (OUT / "quotes").mkdir(parents=True, exist_ok=True)
    make_boq(OUT / "boq-r0.xlsx")
    make_boq_r1(OUT / "boq-r1.xlsx")
    make_vendor_a(OUT / "quotes" / "vendor-a-floors-glass.xlsx")
    make_vendor_b(OUT / "quotes" / "vendor-b-ceilings-partitions.xlsx")
    make_in_house(OUT / "quotes" / "in-house-rate-list.xlsx")
    print(f"demo tender written to {OUT}")


if __name__ == "__main__":
    main()
